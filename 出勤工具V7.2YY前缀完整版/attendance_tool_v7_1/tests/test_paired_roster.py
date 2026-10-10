import io
import unittest
import cv2
import numpy as np
import pandas as pd
from pathlib import Path
from unittest.mock import patch
from streamlit.testing.v1 import AppTest
from openpyxl import load_workbook
from attendance_tool_v7_1.utils.paired_roster import read_id_list, pair_rosters, remap_result
from attendance_tool_v7_1.core.batch import recognize_images, crop_image
from attendance_tool_v7_1.core.export import make_excel


def upload(name, data):
    result = io.BytesIO(data)
    result.name = name
    return result


class PairedRosterTests(unittest.TestCase):
    def test_txt_csv_xlsx_preserve_order_and_text_ids(self):
        expected = ['001', '甲']
        xlsx = io.BytesIO()
        pd.DataFrame(['游戏ID'] + expected).to_excel(xlsx, header=False, index=False)
        for file in [upload('名单.txt', '游戏ID\n001\n甲\n'.encode('utf-8-sig')),
                     upload('名单.csv', '游戏ID\n001\n甲\n'.encode()), upload('名单.xlsx', xlsx.getvalue())]:
            self.assertEqual(read_id_list(file, '游戏'), expected)

    def test_invalid_rows_do_not_silently_shift_pairing(self):
        for content in ['', '甲\n\n乙', '甲\n甲', '游戏ID']:
            with self.assertRaises(ValueError):
                read_id_list(upload('名单.txt', content.encode()), '游戏')
        with self.assertRaises(ValueError):
            read_id_list(upload('名单.csv', b'A,B\n'), 'YY')
        with self.assertRaises(ValueError):
            pair_rosters(['甲', '乙'], ['语音甲'])

    def test_source_ids_map_to_same_person_without_mutating_evidence(self):
        pairs = pair_rosters(['甲', '乙'], ['乙', '甲'])
        aliases = {row['YYID']: row['游戏ID'] for row in pairs}
        region = dict(matched='乙', suggestions=[{'标准姓名': '甲'}], crop=b'original', texts=['乙'])
        original = ({'乙'}, [{'标准姓名': '甲', '区域': 1}], ['乙'], [region])
        found, uncertain, texts, regions = remap_result(original, aliases)
        self.assertEqual(found, {'甲'})
        self.assertEqual(uncertain[0]['标准姓名'], '乙')
        self.assertEqual(regions[0]['matched_id'], '乙')
        self.assertEqual(regions[0]['suggestions'][0]['识别ID'], '甲')
        self.assertEqual(regions[0]['crop'], b'original')
        self.assertEqual(texts, ['乙'])
        self.assertEqual(region['matched'], '乙')
        self.assertNotIn('识别ID', region['suggestions'][0])

    def test_compressed_images_and_crops_are_lossless(self):
        original = np.arange(360, dtype=np.uint8).reshape(10, 12, 3)
        ok, encoded = cv2.imencode('.png', original)
        self.assertTrue(ok)
        def recognize(image, *args):
            np.testing.assert_array_equal(image, original)
            return {'甲'}, [], ['甲'], [dict(id=1, matched='甲', crop=image.copy(), suggestions=[])]
        result = recognize_images([dict(data=encoded.tobytes(), name='test.png', key='1')], '游戏', ['甲'],
                                  recognizer=recognize, compress_crops=True)
        self.assertIsInstance(result[3][0]['crop'], bytes)
        np.testing.assert_array_equal(crop_image(result[3][0]), original)

    def test_export_keeps_both_ids_and_status(self):
        df = pd.DataFrame([{'序号': 1, '游戏 ID': '甲', 'YY ID': '语音甲', '游戏': '☑', 'YY': '×', '总次数': 1, '出勤': '×'}])
        ws = load_workbook(make_excel(df)).active
        self.assertEqual([c.value for c in ws[2]], [1, '甲', '语音甲', '☑', '×', 1, '×'])
        self.assertEqual(ws.column_dimensions['C'].width, 26)

    def test_yy_candidate_manual_confirmation_uses_paired_game_identity(self):
        ok, image = cv2.imencode('.png', np.zeros((10, 12, 3), np.uint8))
        self.assertTrue(ok)
        files = {'names_game': upload('game.txt', '游戏甲\n'.encode()),
                 'names_yy': upload('yy.txt', '语音甲\n'.encode()),
                 'game': [], 'yy': [upload('yy.png', image.tobytes())]}
        candidate = {'标准姓名': '语音甲', '区域': 1, '最高相似度': 80}
        def recognize(image, source, names, *args):
            self.assertEqual(names, ['语音甲'])
            return set(), [candidate], ['近似语音甲'], [dict(id=1, matched=None, reliable=True,
                crop=image.copy(), suggestions=[candidate], texts=['近似语音甲'], error=None,
                method='测试', box=(0,0,12,10))]
        with patch('streamlit.file_uploader', side_effect=lambda *a, **k: files[k['key']]), \
             patch('attendance_tool_v7_1.core.batch.recognize_regions', side_effect=recognize):
            at = AppTest.from_file(str(Path(__file__).parents[1] / 'app.py')).run(timeout=30)
            next(b for b in at.button if '开始识别' in b.label).click().run(timeout=30)
            self.assertFalse(at.exception)
            self.assertEqual(at.session_state['yy_found'], set())
            self.assertEqual(at.session_state['yy_regions'][0]['suggestions'][0]['标准姓名'], '游戏甲')
            at.checkbox(key='review_yy_1_confirmed').check().run(timeout=30)
            self.assertFalse(at.exception)
            self.assertEqual(at.session_state['review_regions'][('yy', 1)], '游戏甲')
            at.checkbox(key='review_yy_1_confirmed').uncheck().run(timeout=30)
            self.assertFalse(at.exception)
            self.assertNotIn(('yy', 1), at.session_state['review_regions'])


if __name__ == '__main__':
    unittest.main()
